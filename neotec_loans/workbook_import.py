"""Neotec Loans — import the client's loan workbook (the IRSAA LOAN layout).

The sheet holds three blocks, found by content rather than fixed rows:
  • term loan     — rows with a due date (C), principal (D) and interest (E);
  • short-term    — rows below the "Financing Date / Payment Date" header:
                    drawdown (C), maturity (D), principal (E), profit (F);
  • charges       — rows labelled "…Charges" with an amount (F).
Header rows that carry dates in the monthly columns (H onwards) map each
column to a month; the monthly cells are the sheet's own accrual.

Options are taken from the sheet where it shows them: a tranche whose first
monthly figure falls after its drawdown month is imported with "Month after
drawdown"; a charge's months and start come from its monthly cells. A charge
row that equals another divided by (1 + VAT) is the same charge net of VAT:
the gross row is imported once as "Amount includes VAT".

Every facility is then cross-checked: the plan computed by loan_core must
reproduce the sheet's monthly figures.
"""
from __future__ import annotations

import datetime as dt
import io

from neotec_loans import loan_core as core

TOL = 0.05


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    return v if isinstance(v, dt.date) else None


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _text(v):
    return str(v).strip() if v not in (None, "") else ""


def read_rows(content: bytes):
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]


def parse(rows, vat_rate: float = 15.0) -> dict:
    facilities, warnings = [], []
    month_cols: dict[int, dt.date] = {}
    block = "term"
    term: dict[str, dict] = {}
    charges = []

    def monthly(r):
        out = {}
        for ci, me in month_cols.items():
            v = _num(r[ci]) if ci < len(r) else None
            if v and abs(v) > 0.005:
                out[me] = out.get(me, 0.0) + v
        return out

    for i, r in enumerate(rows, 1):
        r = list(r) + [None] * 8
        # a header row: dates across the monthly columns (H = index 7 onwards)
        dates = {ci: _date(v) for ci, v in enumerate(r) if ci >= 7 and _date(v)}
        if len(dates) >= 3:
            month_cols = {ci: core.month_end(dv) for ci, dv in dates.items()}
            if "financ" in _text(r[2]).lower():
                block = "short"
            continue
        label, lender = _text(r[0]), _text(r[1])
        if "charge" in label.lower() and _num(r[5]):
            charges.append({"row": i, "lender": lender or "Lender", "amount": _num(r[5]), "months": monthly(r)})
            continue
        if block == "term" and _date(r[2]) and _num(r[3]) is not None and _num(r[4]) is not None and lender:
            t = term.setdefault(lender, {"rows": [], "months": {}, "first_row": i})
            t["rows"].append({"due_date": _date(r[2]), "principal": _num(r[3]), "profit": round(_num(r[4]), 2)})
            for me, v in monthly(r).items():
                t["months"][me] = t["months"].get(me, 0.0) + v
        elif block == "short" and _date(r[2]) and _date(r[3]) and _num(r[4]) and _num(r[5]) is not None:
            start, due = _date(r[2]), _date(r[3])
            m = monthly(r)
            first = min(m) if m else core.month_end(start)
            facilities.append({
                "facility_name": f"{lender or 'Lender'} — {start:%d %b %Y} — {_num(r[4]):,.0f}",
                "facility_type": "Short-term Financing", "lender": lender, "disbursement_date": start,
                "principal": _num(r[4]), "profit_treatment": core.DEFERRED,
                "accrual_basis": core.EQUAL,
                "first_month": core.AFTER if first > core.month_end(start) else "Month of drawdown (or previous due date)",
                "schedule": [{"due_date": due, "principal": _num(r[4]), "profit": round(_num(r[5]), 2)}],
                "sheet_months": m, "source_row": i})

    for lender, t in term.items():
        first_month = min(t["months"]) if t["months"] else core.month_end(t["rows"][0]["due_date"])
        facilities.insert(0, {
            "facility_name": f"{lender} — Term Loan", "facility_type": "Term Loan", "lender": lender,
            "disbursement_date": first_month.replace(day=1), "principal": sum(x["principal"] for x in t["rows"]),
            "profit_treatment": core.ACCRUE, "accrual_basis": core.EQUAL,
            "first_month": "Month of drawdown (or previous due date)",
            "schedule": t["rows"], "sheet_months": t["months"], "source_row": t["first_row"]})

    # charges: a row equal to another / (1 + VAT) is that charge net of VAT
    rate = vat_rate / 100.0
    used = set()
    for a in charges:
        for b in charges:
            if a is not b and a["lender"] == b["lender"] and abs(a["amount"] / (1 + rate) - b["amount"]) < 1:
                used.add(id(b))
                a["net_twin"] = b
                warnings.append(f"Rows {a['row']} and {b['row']}: the same {a['lender']} charge, gross and net of VAT "
                                f"— imported once ({a['amount']:,.2f} including VAT).")
    for c in charges:
        if id(c) in used:
            continue
        twin = c.get("net_twin")
        months = (twin or c)["months"]
        fee_vat = "Amount includes VAT" if twin else "No VAT"
        if not twin:
            warnings.append(f"Row {c['row']}: {c['lender']} charge {c['amount']:,.2f} imported as “No VAT” because the "
                            f"sheet amortises the full amount; change it if the charge includes VAT.")
        start = min(months) if months else None
        facilities.append({
            "facility_name": f"{c['lender']} — Facility charge {c['amount']:,.0f}", "facility_type": "Facility Charge",
            "lender": c["lender"], "principal": 0, "fee_amount": c["amount"], "fee_vat": fee_vat, "vat_rate": vat_rate,
            "fee_months": len(months), "fee_start_date": start.replace(day=1) if start else None,
            "schedule": [], "sheet_months": months, "source_row": c["row"]})
        warnings.append(f"Row {c['row']}: set “Charge paid on” on “{c['lender']} — Facility charge "
                        f"{c['amount']:,.0f}” before posting its payment.")

    # unique names
    seen: dict[str, int] = {}
    for f in facilities:
        n = f["facility_name"]
        seen[n] = seen.get(n, 0) + 1
        if seen[n] > 1:
            f["facility_name"] = f"{n} #{seen[n]}"
    for f in facilities:
        f["check"] = cross_check(f)
    return {"facilities": facilities, "warnings": warnings}


ROUNDING = 0.25


def cross_check(f: dict) -> dict:
    """Compare the computed monthly plan with the sheet's monthly figures.

    status: "exact"     every month within 0.01;
            "rounding"  every month within 0.25 and the computed total equals
                        the bank's figures (schedule profit / charge) — the
                        sheet's own monthly rounding drifts;
            "different" anything else: review before posting."""
    plan = core.fee_plan(f) if f["facility_type"] == "Facility Charge" else core.profit_plan(f)
    sheet = f.get("sheet_months") or {}
    diffs = []
    worst = 0.0
    for me in sorted(set(plan) | set(sheet)):
        a, b = plan.get(me, 0.0), sheet.get(me, 0.0)
        worst = max(worst, abs(a - b))
        if abs(a - b) > 0.01:
            diffs.append({"month": me.isoformat(), "computed": round(a, 2), "sheet": round(b, 2)})
    if f["facility_type"] == "Facility Charge":
        target = core.fee_net(f)[0]
    else:
        target = sum(float(r.get("profit") or 0) for r in f.get("schedule") or [])
    total_ok = abs(sum(plan.values()) - target) < 0.01
    status = "exact" if not diffs else ("rounding" if worst <= ROUNDING and total_ok else "different")
    note = ""
    if status == "rounding":
        note = (f"Matches the bank's figures ({target:,.2f}); the sheet's monthly rounding differs by up to "
                f"{worst:.2f} a month (sheet total {sum(sheet.values()):,.2f}).")
    return {"ok": status != "different", "status": status, "months": len(plan), "worst": round(worst, 4),
            "note": note, "differences": diffs[:12]}
