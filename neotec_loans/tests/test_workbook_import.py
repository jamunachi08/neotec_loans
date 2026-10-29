"""Workbook importer — on a synthetic workbook in the client's exact layout
(the real file holds the company's loan figures and is not shipped)."""
import datetime as dt
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from neotec_loans import loan_core as c  # noqa: E402
from neotec_loans import workbook_import as wi  # noqa: E402

D = dt.datetime


def build() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    months = [D(2025, 12, 25), D(2026, 1, 26), D(2026, 2, 26), D(2026, 3, 26), D(2026, 4, 26), D(2026, 5, 26),
              D(2026, 6, 26)]
    ws.cell(row=2, column=3, value="Date"); ws.cell(row=2, column=4, value="Amount")
    for i, m in enumerate(months):
        ws.cell(row=2, column=8 + i, value=m)
    # term loan: 2 instalments; monthly cells as the sheet lays them out
    for r, (due, pr, it, cells) in enumerate([
            (D(2026, 3, 1), 37053, 16427.26, [5475.7533] * 3),
            (D(2026, 6, 1), 37053, 14582.83, [0, 0, 0, 4860.9433, 4860.9433, 4860.9434])], start=3):
        ws.cell(row=r, column=1, value="Intrest"); ws.cell(row=r, column=2, value="BankA")
        ws.cell(row=r, column=3, value=due); ws.cell(row=r, column=4, value=pr); ws.cell(row=r, column=5, value=it)
        for i, v in enumerate(cells):
            ws.cell(row=r, column=8 + i, value=v)
    ws.cell(row=5, column=1, value="Intrest"); ws.cell(row=5, column=4, value=74106)        # totals row: ignored
    ws.cell(row=7, column=6, value=0.0401)                                               # helper cell: ignored
    # short-term block
    ws.cell(row=9, column=3, value="Financing \nDate "); ws.cell(row=9, column=4, value="Payment \nDate")
    for i, m in enumerate(months):
        ws.cell(row=9, column=8 + i, value=m)
    for r, (start, due, pr, pf, cells) in enumerate([
            (D(2025, 11, 19), D(2026, 2, 17), 599841, 23987.64, [7995.88] * 3),           # starts month after
            (D(2026, 1, 6), D(2026, 4, 6), 592365, 23749.73, [0, 7916.58, 7916.58, 7916.57])], start=10):  # same month
        ws.cell(row=r, column=1, value="Intrest"); ws.cell(row=r, column=2, value="BankB")
        ws.cell(row=r, column=3, value=start); ws.cell(row=r, column=4, value=due)
        ws.cell(row=r, column=5, value=pr); ws.cell(row=r, column=6, value=pf)
        for i, v in enumerate(cells):
            ws.cell(row=r, column=8 + i, value=v)
    # charges: gross + net twin (one charge), and a single charge
    for r, (lender, amt, cells) in enumerate([("BankB", 23000, [0, 11500.0, 11500.0]),
                                              ("BankB", 20000, [0, 10000.0, 10000.0]),
                                              ("BankC", 1200, [600.0, 600.0])], start=14):
        ws.cell(row=r, column=1, value="Facilities Charges"); ws.cell(row=r, column=2, value=lender)
        ws.cell(row=r, column=6, value=amt)
        for i, v in enumerate(cells):
            ws.cell(row=r, column=8 + i, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Importer(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res = wi.parse(wi.read_rows(build()))
        cls.by = {f["facility_name"]: f for f in cls.res["facilities"]}

    def test_blocks_found(self):
        types = [f["facility_type"] for f in self.res["facilities"]]
        self.assertEqual(types.count("Term Loan"), 1)
        self.assertEqual(types.count("Short-term Financing"), 2)
        self.assertEqual(types.count("Facility Charge"), 2)                 # the net twin is not a second charge

    def test_term_loan(self):
        t = self.by["BankA — Term Loan"]
        self.assertEqual(len(t["schedule"]), 2)
        self.assertEqual(t["principal"], 74106)
        self.assertEqual(t["profit_treatment"], c.ACCRUE)
        self.assertIn(t["check"]["status"], ("exact", "rounding"))         # sheet cells carry 4 decimals

    def test_first_month_read_from_the_sheet(self):
        nov = self.by["BankB — 19 Nov 2025 — 599,841"]
        jan = self.by["BankB — 06 Jan 2026 — 592,365"]
        self.assertEqual(nov["first_month"], c.AFTER)
        self.assertNotEqual(jan["first_month"], c.AFTER)
        self.assertEqual((nov["check"]["status"], jan["check"]["status"]), ("exact", "exact"))

    def test_charges(self):
        g = self.by["BankB — Facility charge 23,000"]
        self.assertEqual((g["fee_vat"], g["fee_months"], g["fee_start_date"]), ("Amount includes VAT", 2, dt.date(2026, 1, 1)))
        self.assertEqual(g["check"]["status"], "exact")                      # 20,000 net over 2 months = the net row
        single = self.by["BankC — Facility charge 1,200"]
        self.assertEqual((single["fee_vat"], single["fee_months"]), ("No VAT", 2))
        joined = " ".join(self.res["warnings"])
        self.assertIn("gross and net of VAT", joined)
        self.assertIn("imported as “No VAT”", joined)

    def test_a_real_difference_is_reported(self):
        f = dict(self.by["BankB — 06 Jan 2026 — 592,365"])
        f["sheet_months"] = {k: v + 50 for k, v in f["sheet_months"].items()}
        self.assertEqual(wi.cross_check(f)["status"], "different")


class SheetRounding(unittest.TestCase):
    def test_small_monthly_drift_with_bank_totals_is_rounding(self):
        f = {"facility_type": "Short-term Financing", "disbursement_date": dt.date(2026, 1, 6), "first_month": "",
             "schedule": [{"due_date": dt.date(2026, 4, 6), "principal": 1, "profit": 300.0}],
             "sheet_months": {dt.date(2026, 1, 31): 100.1, dt.date(2026, 2, 28): 99.9, dt.date(2026, 3, 31): 100.1}}
        ck = wi.cross_check(f)
        self.assertEqual(ck["status"], "rounding")
        self.assertIn("300.00", ck["note"])


class Endpoint(unittest.TestCase):
    def test_import_endpoint_is_post_only_and_checks_create_permission(self):
        src = (Path(__file__).resolve().parents[1] / "api.py").read_text()
        i = src.index("def import_workbook(")
        self.assertIn('@frappe.whitelist(methods=["POST"])', src[i - 60:i])
        self.assertIn('frappe.has_permission(DT, "create")', src[i:i + 900])


if __name__ == "__main__":
    unittest.main()
